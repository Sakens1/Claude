#!/usr/bin/env python3
"""Genera la solución desarrollada del Taller 2 (energía en movimiento de rototraslación) sobre la
plantilla original, con datos de ejemplo o con los datos de un laboratorio.

Sin argumentos usa datos de ejemplo (anillo de hierro en un plano levemente inclinado).
Con datos propios acepta las mismas opciones que verificar_energia.py, más --etiqueta y --salida.

    python3 generar_solucion.py
    python3 generar_solucion.py --M 0.45 --Rint 0.025 --Rext 0.035 --h0 0.06 --largo 1.2 --D 1.2 \
        --d 0 0.01 ... --t 0 0.25 ... --etiqueta "Lab jueves" --salida ../solucion/lab_jueves.docx

Requiere python-docx y matplotlib.
"""
import math
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
import docx
from docx.oxml.ns import qn
from docx.table import Table

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(RAIZ, "..", "comun"))
import verificar_energia as ve  # noqa: E402
from docx_util import coma, cient, llenar_cuadro, agregar_a_celda, celda_simple  # noqa: E402

ENUNCIADO = os.path.join(RAIZ, "fuentes", "Taller2_Rototraslacion_IN1090C-IN1111C_enunciado.docx")
SALIDA_POR_DEFECTO = os.path.join(RAIZ, "solucion", "Taller2_Rototraslacion_RESUELTO.docx")
FILAS = 11  # filas de datos de las Tablas 2 y 3 de la plantilla

# Datos de ejemplo: anillo de hierro (M = 450 g, Rint = 25 mm, Rext = 35 mm) que parte del reposo
# en lo alto de un plano de 1,20 m elevado 6,0 cm; posiciones registradas cada 0,25 s.
EJEMPLO = ["--M", "0.450", "--Rint", "0.025", "--Rext", "0.035", "--h0", "0.060", "--largo", "1.20",
           "--D", "1.20",
           "--d", "0", "0.010", "0.030", "0.075", "0.130", "0.215", "0.305", "0.420", "0.540", "0.685", "0.845",
           "--t", "0", "0.25", "0.50", "0.75", "1.00", "1.25", "1.50", "1.75", "2.00", "2.25", "2.50"]


def sg(x, dec):
    return (" + " if x >= 0 else " − ") + coma(abs(x), dec)


def fmt_eje(v, _):
    return "0" if v == 0 else f"{v:.3f}".rstrip("0").rstrip(".").replace(".", ",")


def mt(x, dec):
    return coma(x, dec).replace(",", "{,}")


def mts(x, dec):
    return ("+" if x >= 0 else "-") + mt(abs(x), dec)


# ---------------------------------------------------------------- gráficos
def graficos(r, base):
    filas = r["tabla3"]
    d = [f["d"] for f in filas]
    t = [f["t"] for f in filas]
    A, B, C = r["ajuste_d_t"]["A"], r["ajuste_d_t"]["B"], r["ajuste_d_t"]["C"]
    rutas = {}

    fig, ax = plt.subplots(figsize=(7, 4.2), dpi=200)
    ax.scatter(t, d, color="#1f4e79", zorder=3, label="Datos (Tabla 2)")
    ts = [max(t) * k / 100 for k in range(101)]
    ax.plot(ts, [A * x * x + B * x + C for x in ts], color="#c00000", lw=1.4, label="Ajuste polinómico de grado 2")
    ax.set_xlabel("Tiempo  t [s]")
    ax.set_ylabel("Posición  d [m]")
    ax.set_title("Posición del anillo en función del tiempo")
    ax.text(0.04, 0.95, f"$d(t) = {mt(A, 4)}\\,t^2 {mts(B, 4)}\\,t {mts(C, 4)}$\n$R^2 = {mt(r['ajuste_d_t']['R2'], 4)}$",
            transform=ax.transAxes, va="top", fontsize=10.5, bbox=dict(boxstyle="round", fc="white", ec="#888"))
    for a_ in (ax.xaxis, ax.yaxis):
        a_.set_major_formatter(FuncFormatter(fmt_eje))
    ax.grid(True, ls=":", alpha=0.6)
    ax.legend(loc="lower right", fontsize=9)
    fig.tight_layout()
    rutas["dt"] = base + "_d_vs_t.png"
    fig.savefig(rutas["dt"])
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 4.2), dpi=200)
    ax.plot(d, [f["Ktras"] for f in filas], "o-", color="#1f4e79", label="$K_{Tras}$")
    ax.plot(d, [f["Krot"] for f in filas], "s-", color="#e07b00", label="$K_{Rot}$")
    ax.set_xlabel("Posición  d [m]")
    ax.set_ylabel("Energía [J]")
    ax.set_title("Energía cinética de traslación y de rotación vs posición")
    for a_ in (ax.xaxis, ax.yaxis):
        a_.set_major_formatter(FuncFormatter(fmt_eje))
    ax.grid(True, ls=":", alpha=0.6)
    ax.legend(loc="upper left", fontsize=9)
    fig.tight_layout()
    rutas["k"] = base + "_Krot_Ktras.png"
    fig.savefig(rutas["k"])
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7, 4.9), dpi=200)
    series = [("Ktot", "$K_{Tot}$", "#1f4e79", "o"), ("Ug", "$U_g$", "#2e7d32", "s"), ("Emec", "$E_{mec}$", "#c00000", "^")]
    xs = [min(d), max(d)]
    texto = []
    for clave, etq, col, mk in series:
        ys = [f[clave] for f in filas]
        m, b, rr = r["ajustes_lineales"][clave]
        ax.scatter(d, ys, color=col, marker=mk, zorder=3, label=etq)
        ax.plot(xs, [m * x + b for x in xs], color=col, lw=1.1, ls="--")
        texto.append(f"{etq}$ = {mt(m, 4)}\\,d {mts(b, 4)}$")
    ax.set_xlabel("Posición  d [m]")
    ax.set_ylabel("Energía [J]")
    ax.set_title("Energía cinética total, potencial y mecánica vs posición")
    fig.text(0.5, 0.015, "Ajustes lineales:   " + "     ".join(texto), ha="center", va="bottom", fontsize=9)
    for a_ in (ax.xaxis, ax.yaxis):
        a_.set_major_formatter(FuncFormatter(fmt_eje))
    ax.set_ylim(0, max(f["Emec"] for f in filas) * 1.15)
    ax.grid(True, ls=":", alpha=0.6)
    ax.legend(loc="upper left", fontsize=9, ncol=3)
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    rutas["e"] = base + "_energias.png"
    fig.savefig(rutas["e"])
    plt.close(fig)
    return rutas


# ---------------------------------------------------------------- respuestas
def respuestas(r, rutas):
    e = r["entradas_SI"]
    aj = r["ajuste_d_t"]
    A, B, C = aj["A"], aj["B"], aj["C"]
    beta = r["I_sobre_MR2"]
    frac_rot = beta / (1 + beta) * 100
    Em, s = r["Emec_promedio"], r["sigma"]
    srel = s / Em * 100
    var = r["variacion_Emec_pct"]
    mK, bK, _ = r["ajustes_lineales"]["Ktot"]
    mU, bU, _ = r["ajustes_lineales"]["Ug"]
    mE, bE, rE = r["ajustes_lineales"]["Emec"]
    mi = r["mitad_recorrido"]
    a_exp, a_teo = r["aceleracion_experimental"], r["aceleracion_teorica"]
    dif_a = (a_exp - a_teo) / a_teo * 100
    filas = r["tabla3"]
    ultimo = filas[-1]
    conserva = srel <= 5

    R = {}
    R["dt"] = [
        ("b", "Gráfico"),
        ("img", rutas["dt"]),
        ("p", "Línea de tendencia: ajuste polinómico de grado 2 (movimiento con aceleración constante)."),
        ("eq", f"**d(t) = {coma(A, 4)}·t^{{2}}{sg(B, 4)}·t{sg(C, 4)}   [m]**     R^{{2}} = {coma(aj['R2'], 4)}"),
        ("p", "Derivando respecto del tiempo se obtiene la rapidez del centro de masa:"),
        ("eq", f"**v(t) = d'(t) = {coma(2 * A, 4)}·t{sg(B, 4)}   [m/s]**"),
        ("p", f"La aceleración del centro de masa es a = 2·{coma(A, 4)} = {coma(a_exp, 3)} m/s^{{2}}. "
              f"Para un anillo que rueda sin deslizar se espera a = g·sen θ/(1 + I/(M·R^{{2}})) = {coma(a_teo, 3)} m/s^{{2}} "
              f"(diferencia de {coma(dif_a, 1)} %)."),
    ]
    R["hip"] = [
        ("p", "Si el anillo parte del reposo y rueda sin deslizar sobre el plano inclinado, la fuerza de roce "
              "estático no realiza trabajo y la única fuerza que trabaja es el peso (conservativa). Por lo tanto, "
              "a medida que el anillo desciende (aumenta d), su energía potencial gravitatoria U_{g} = M·g·h "
              "disminuirá linealmente con la posición, mientras que sus energías cinéticas de traslación "
              "(½·M·v^{2}) y de rotación (½·I·ω^{2}) aumentarán en la misma cantidad, de modo que la energía "
              "mecánica E_{mec} = K_{Tot} + U_{g} se mantendrá constante."),
        ("p", f"Además, como v = ω·R, la razón K_{{Rot}}/K_{{Tras}} = I/(M·R^{{2}}) es constante "
              f"(≈ {coma(beta, 2)} para este anillo): cerca del {coma(frac_rot, 0)} % de la energía cinética "
              "corresponderá a rotación, por lo que el anillo no puede modelarse como partícula."),
    ]
    R["k"] = [
        ("img", rutas["k"]),
        ("p", "**Inferencia:** ambas energías parten de cero y aumentan a medida que el anillo avanza, de forma "
              "aproximadamente lineal con la posición (con aceleración constante, v^{2} crece proporcional a d). "
              f"La energía de traslación es siempre mayor que la de rotación y la razón entre ellas se mantiene "
              f"constante: K_{{Rot}}/K_{{Tras}} ≈ {coma(ultimo['Krot'] / ultimo['Ktras'], 2)}, igual a "
              f"I/(M·R^{{2}}) = {coma(beta, 2)}. Es decir, del total de energía cinética, cerca del "
              f"{coma(frac_rot, 0)} % se usa en hacer girar el anillo. Esto muestra que, para un cuerpo con "
              "momento de inercia no despreciable, el modelo de partícula subestima la energía cinética y es "
              "necesario el modelo de cuerpo rígido."),
    ]
    R["e"] = [
        ("img", rutas["e"]),
        ("p", "Ajustes lineales de los datos (energías en J, d en m):"),
        ("eq", f"K_{{Tot}} = {coma(mK, 4)}·d{sg(bK, 4)}     U_{{g}} = {coma(mU, 4)}·d{sg(bU, 4)}     "
               f"E_{{mec}} = {coma(mE, 4)}·d{sg(bE, 4)}"),
        ("p", "La energía cinética total aumenta y la potencial disminuye con pendientes de igual magnitud y "
              "signo contrario, por lo que su suma, la energía mecánica, es prácticamente horizontal "
              f"(pendiente {coma(mE, 4)} J/m, casi nula)."),
        ("lista", f"**Energía mecánica total del sistema:** E_{{mec}} ≈ {coma(Em, 3)} J (promedio de la Tabla 3; "
                  f"igual a la energía potencial inicial M·g·h_{{0}} = {coma(filas[0]['Ug'], 3)} J)."),
        ("lista", f"**A la mitad del recorrido** (d = {coma(mi['d'], 3)} m), según los ajustes: "
                  f"K_{{Tot}} ≈ {coma(mi['Ktot'], 3)} J y U_{{g}} ≈ {coma(mi['Ug'], 3)} J."),
        ("p", f"La energía mecánica disminuye levemente ({coma(abs(var), 1)} % entre el máximo y el mínimo), "
              "lo que se atribuye a errores al leer las posiciones y tiempos y a pequeñas pérdidas por roce de "
              "rodadura; dentro de esa precisión, la energía mecánica se conserva."),
    ]
    R["prom"] = [
        ("eq", f"<E_{{Tot}}> = (1/N)·Σ E_{{Tot,i}} = {coma(Em, 4)} J"),
        ("eq", f"ΔE_{{Tot}} = σ = √( Σ (E_{{Tot,i}} − <E_{{Tot}}>)^{{2}} / N ) = {coma(s, 4)} J   (N = {len(filas)})"),
        ("eq", f"**E_{{Tot}} = ({coma(Em, 3)} ± {coma(s, 3)}) J**"),
        ("p", f"El error relativo es σ/<E> = {coma(srel, 1)} %, "
              + ("menor al 5 %: la energía mecánica se mantiene prácticamente constante a lo largo del recorrido."
                 if conserva else "mayor al 5 %: hay una variación apreciable de la energía mecánica, que debe "
                 "explicarse por errores de medición o pérdidas.")),
    ]
    R["concl"] = [
        ("p", "Respecto del primer resultado de aprendizaje, se analizó gráficamente cómo cambian las energías "
              "que forman la energía mecánica del anillo: al descender, la energía potencial gravitatoria disminuye "
              "linealmente con la posición, mientras que las energías cinéticas de traslación y de rotación aumentan. "
              f"La energía de rotación representa cerca del {coma(frac_rot, 0)} % de la energía cinética, lo que "
              "confirma que el anillo debe tratarse como cuerpo rígido y que su momento de inercia "
              f"(I = {cient(r['I'], 2)} kg·m^{{2}}) es indispensable para calcular su energía."),
        ("p", "Respecto del segundo resultado de aprendizaje, la energía mecánica se mantuvo prácticamente "
              f"constante: E_{{Tot}} = ({coma(Em, 3)} ± {coma(s, 3)}) J, con una dispersión de {coma(srel, 1)} % y una "
              "línea de tendencia de pendiente casi nula. Por lo tanto, "
              + ("la hipótesis se confirma: " if conserva else "la hipótesis se confirma solo parcialmente: ")
              + "al rodar sin deslizar, el roce estático no realiza trabajo y la energía mecánica se conserva, "
              "transformándose energía potencial en energía cinética de traslación y de rotación."),
        ("p", f"La aceleración obtenida del ajuste ({coma(a_exp, 3)} m/s^{{2}}) es {coma(abs(dif_a), 1)} % "
              f"{'menor' if dif_a < 0 else 'mayor'} que la esperada para rodadura sin deslizamiento "
              f"({coma(a_teo, 3)} m/s^{{2}}). Las diferencias se explican por la lectura de posiciones y tiempos "
              "(resolución de la escala y del intervalo de tiempo) y por pequeñas pérdidas de energía. Para mejorar, "
              "se pueden registrar más puntos y con intervalos de tiempo más cortos."),
    ]
    return R


def decimales(valores, minimo, maximo):
    for k in range(minimo, maximo + 1):
        if all(abs(round(v, k) - v) < 1e-12 for v in valores):
            return k
    return maximo


def main(argv=None):
    p = ve.construir_parser()
    p.add_argument("--etiqueta", default=None)
    p.add_argument("--salida", default=SALIDA_POR_DEFECTO)
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv or all(x.startswith("--etiqueta") or x.startswith("--salida") for x in argv[::2]):
        argv = EJEMPLO + argv
        etiqueta_def = "Datos de ejemplo (anillo de hierro)"
    else:
        etiqueta_def = ""
    a = p.parse_args(argv)
    r = ve.analizar(a)
    if len(r["tabla3"]) > FILAS:
        p.error(f"la plantilla tiene {FILAS} filas en las Tablas 2 y 3; se recibieron {len(r['tabla3'])} mediciones")

    salida = os.path.abspath(a.salida)
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    base = os.path.join(os.path.dirname(salida), "grafico") if salida == os.path.abspath(SALIDA_POR_DEFECTO) \
        else os.path.splitext(salida)[0]
    rutas = graficos(r, base)
    R = respuestas(r, rutas)

    d = docx.Document(ENUNCIADO)
    cuerpo = list(d.element.body.iterchildren())

    def tabla(i, filas=None, cols=None):
        el = cuerpo[i]
        assert el.tag == qn("w:tbl"), f"elemento {i} no es tabla"
        T = Table(el, d)
        if filas:
            assert len(T.rows) == filas and len(T.columns) == cols, f"tabla {i} con forma inesperada"
        return T

    # título: agrega la etiqueta (laboratorio / ejemplo)
    etq = a.etiqueta or etiqueta_def
    if etq:
        agregar_a_celda(tabla(0).cell(0, 0), [("p", f"SOLUCIÓN DESARROLLADA / REFERENCIA DOCENTE: {etq}")])

    e = r["entradas_SI"]
    t1 = tabla(10, 3, 5)
    for j, txt in enumerate([coma(e["M"], 3), coma(e["Rint"], 4), coma(e["Rext"], 4),
                             cient(r["I"], 3), coma(e["theta_grados"], 2)]):
        celda_simple(t1.cell(1, j), txt)
    celda_simple(t1.cell(2, 4), f"sen θ = h_{{0}}/L = {coma(e['sen_theta'], 4)}", tam=8)

    filas = r["tabla3"]
    dd = decimales([f["d"] for f in filas], 3, 4)
    dt = decimales([f["t"] for f in filas], 2, 3)
    t2 = tabla(16, FILAS + 1, 2)
    t3 = tabla(33, FILAS + 1, 9)
    for i, f in enumerate(filas):
        celda_simple(t2.cell(i + 1, 0), coma(f["d"], dd))
        celda_simple(t2.cell(i + 1, 1), coma(f["t"], dt))
        vals = [coma(f["d"], dd), coma(f["t"], dt), coma(f["h"], 4), coma(f["v"], 4)] + \
               [coma(f[k], 4) for k in ("Krot", "Ktras", "Ktot", "Ug", "Emec")]
        for j, v in enumerate(vals):
            celda_simple(t3.cell(i + 1, j), v, tam=8)

    llenar_cuadro(tabla(29, 1, 1).cell(0, 0), R["dt"])
    agregar_a_celda(tabla(37, 1, 1).cell(0, 0), R["hip"])
    t40 = tabla(40, 2, 1)
    agregar_a_celda(t40.cell(0, 0), R["k"])
    agregar_a_celda(t40.cell(1, 0), R["e"])
    llenar_cuadro(tabla(48, 1, 1).cell(0, 0), R["prom"])
    llenar_cuadro(tabla(52, 1, 1).cell(0, 0), R["concl"])

    d.save(salida)
    print("Generado:", salida)
    for v in rutas.values():
        print("Generado:", v)
    ve.imprimir(r)


if __name__ == "__main__":
    main()
