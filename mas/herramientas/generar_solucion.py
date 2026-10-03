#!/usr/bin/env python3
"""Genera la solución desarrollada del Taller de Movimiento Armónico Simple (resorte) sobre la plantilla
original, con datos de ejemplo o con los datos de un laboratorio.

Sin argumentos usa datos de ejemplo. Con datos propios acepta las mismas opciones que verificar_mas.py
(deben venir ambas partes), más --etiqueta y --salida:

    python3 generar_solucion.py
    python3 generar_solucion.py --masas-en-g --x-en-mm --m1 50 100 ... --x 20 39 ... \
        --m2 50 100 ... --t10 3.16 4.11 ... --etiqueta "Lab lunes" --salida ../solucion/lab_lunes.docx

Requiere python-docx y matplotlib.
"""
import math
import os
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
import docx
from docx.oxml.ns import qn
from docx.table import Table

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.join(RAIZ, "..", "comun"))
import verificar_mas as vm  # noqa: E402
from docx_util import coma, agregar_a_celda, celda_simple, agregar_texto  # noqa: E402

ENUNCIADO = os.path.join(RAIZ, "fuentes", "Taller_MAS_IN1090C-IN1111C_enunciado.docx")
SALIDA_POR_DEFECTO = os.path.join(RAIZ, "solucion", "Taller_MAS_RESUELTO.docx")
FILAS1, FILAS2 = 6, 8

# Datos de ejemplo: resorte de estiramiento (k ≈ 25 N/m, masa ≈ 30 g), masas de 50 a 400 g,
# deformaciones leídas con regla (mm) y tiempos de 10 oscilaciones con cronómetro.
EJEMPLO = ["--masas-en-g", "--x-en-mm",
           "--m1", "50", "100", "150", "200", "250", "300",
           "--x", "20", "39", "59", "78", "98", "117",
           "--m2", "50", "100", "150", "200", "250", "300", "350", "400",
           "--t10", "3.16", "4.11", "5.13", "5.67", "6.46", "6.93", "7.63", "8.00"]


def sg(x, d):
    return (" + " if x >= 0 else " − ") + coma(abs(x), d)


def mt(x, d):
    return coma(x, d).replace(",", "{,}")


def mts(x, d):
    return ("+" if x >= 0 else "-") + mt(abs(x), d)


def fmt_eje(v, _):
    return "0" if v == 0 else f"{v:.3f}".rstrip("0").rstrip(".").replace(".", ",")


def grafico(xs, ys, m, b, r2, xl, yl, titulo, eq, ruta):
    fig, ax = plt.subplots(figsize=(5.2, 3.9), dpi=200)
    ax.scatter(xs, ys, color="#1f4e79", zorder=3, label="Datos")
    xx = [0, max(xs) * 1.05]
    ax.plot(xx, [m * v + b for v in xx], color="#c00000", lw=1.3, label="Línea de tendencia lineal")
    ax.set_xlabel(xl)
    ax.set_ylabel(yl)
    ax.set_title(titulo, fontsize=10.5)
    ax.text(0.04, 0.95, eq + f"\n$R^2 = {mt(r2, 4)}$", transform=ax.transAxes, va="top", fontsize=9.5,
            bbox=dict(boxstyle="round", fc="white", ec="#888"))
    for a_ in (ax.xaxis, ax.yaxis):
        a_.set_major_formatter(FuncFormatter(fmt_eje))
    ax.set_xlim(0, max(xs) * 1.08)
    ax.set_ylim(0, max(ys) * 1.15)
    ax.grid(True, ls=":", alpha=0.6)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    fig.savefig(ruta)
    plt.close(fig)


def main(argv=None):
    p = vm.construir_parser()
    p.add_argument("--etiqueta", default=None)
    p.add_argument("--salida", default=SALIDA_POR_DEFECTO)
    argv = list(sys.argv[1:] if argv is None else argv)
    etiqueta_def = ""
    if not any(x in argv for x in ("--m1", "--m2")):
        argv = EJEMPLO + argv
        etiqueta_def = "Datos de ejemplo (resorte de estiramiento)"
    a = p.parse_args(argv)
    r = vm.analizar(a)
    if "parte1" not in r or "parte2" not in r:
        p.error("se necesitan los datos de ambas partes")
    p1, p2, cp = r["parte1"], r["parte2"], r["comparacion"]
    if len(p1["x_m"]) > FILAS1 or len(p2["T_s"]) > FILAS2:
        p.error(f"la plantilla tiene {FILAS1} filas en la Tabla 1 y {FILAS2} en la Tabla 2")

    salida = os.path.abspath(a.salida)
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    base = os.path.join(os.path.dirname(salida), "grafico") if salida == os.path.abspath(SALIDA_POR_DEFECTO) \
        else os.path.splitext(salida)[0]

    a1, b1, r21 = (p1["ajuste_x_vs_F"][k] for k in ("pendiente_m_por_N", "intercepto_m", "R2"))
    a2, b2, r22 = (p2["ajuste_T2_vs_m"][k] for k in ("pendiente_s2_por_kg", "intercepto_s2", "R2"))
    k1, k2 = p1["k1_N_por_m"], p2["k2_N_por_m"]
    mef = p2["masa_efectiva_resorte_kg"]
    dif = cp["diferencia_pct_respecto_promedio"]
    g1, g2 = base + "_1_x_vs_F.png", base + "_2_T2_vs_m.png"
    grafico(p1["F_N"], p1["x_m"], a1, b1, r21, "Fuerza aplicada (peso)  F [N]", "Deformación  x [m]",
            "Gráfico 1: Deformación del resorte vs fuerza aplicada",
            f"$x = {mt(a1, 4)}\\,F {mts(b1, 4)}$", g1)
    grafico(p2["masas_kg"], p2["T2_s2"], a2, b2, r22, "Masa colgada  m [kg]", "Periodo al cuadrado  T² [s²]",
            "Gráfico 2: Periodo al cuadrado vs masa",
            f"$T^2 = {mt(a2, 4)}\\,m {mts(b2, 4)}$", g2)

    d = docx.Document(ENUNCIADO)
    cuerpo = list(d.element.body.iterchildren())

    def tabla(i):
        assert cuerpo[i].tag == qn("w:tbl"), f"elemento {i} no es tabla"
        return Table(cuerpo[i], d)

    etq = a.etiqueta or etiqueta_def
    if etq:
        agregar_a_celda(tabla(0).cell(0, 0), [("p", f"SOLUCIÓN DESARROLLADA / REFERENCIA DOCENTE: {etq}")])

    # --- Parte I
    agregar_a_celda(tabla(20).cell(0, 0), [
        ("p", "Al aumentar el peso del cuerpo colgado, la deformación del resorte **aumenta**, porque en el "
              "equilibrio la fuerza que ejerce el resorte iguala al peso (k·x = m·g): a mayor fuerza aplicada, "
              "mayor estiramiento."),
        ("p", "La dependencia será **lineal** (directamente proporcional): según la ley de Hooke x = F/k, con k "
              "constante dentro del rango elástico. El gráfico de deformación versus fuerza debería ser una recta "
              "que pasa aproximadamente por el origen, cuya pendiente es 1/k."),
    ])
    t1 = tabla(24)
    inner1 = t1.cell(0, 0).tables[0]
    for i, (m, F, x) in enumerate(zip(p1["masas_kg"], p1["F_N"], p1["x_m"])):
        fila = inner1.rows[i + 1]
        celda_simple(fila.cells[1], coma(m, 3), tam=9)
        celda_simple(fila.cells[2], coma(F, 3), tam=9)
        celda_simple(fila.cells[3], coma(x, 3), tam=9)
    agregar_a_celda(t1.cell(0, 0), [("p", "F = m·g, con g = 9,8 m/s²; x = deformación medida desde el largo "
                                          "natural del resorte.")])
    agregar_a_celda(t1.cell(0, 1), [("img", g1), ("eq", f"x = {coma(a1, 4)}·F{sg(b1, 4)}   (x en m, F en N)")],
                    ancho_img_cm=8)
    agregar_a_celda(tabla(28).cell(0, 0), [
        ("p", "De la ley de Hooke, x = (1/k)·F: la pendiente del gráfico 1 es 1/k. Por lo tanto:"),
        ("eq", f"k = 1/pendiente = 1/({coma(a1, 5)} m/N) = **{coma(k1, 1)} N/m**"),
    ])
    agregar_a_celda(tabla(31).cell(0, 0), [
        ("p", f"Sí. Los puntos se alinean en una recta (R^{{2}} = {coma(r21, 4)}) con intercepto prácticamente nulo "
              f"({coma(b1 * 1000, 1)} mm), por lo que la deformación es directamente proporcional a la fuerza, "
              "como predice la ley de Hooke y se planteó en la hipótesis."),
        ("eq", f"Relación algebraica obtenida: x = {coma(a1, 4)}·F{sg(b1, 4)}   ⟹   F ≈ {coma(k1, 1)}·x"),
    ])

    # --- Parte II
    agregar_a_celda(tabla(48).cell(0, 0), [
        ("p", "Al aumentar la masa colgada, el periodo de oscilación **aumenta**: con más masa (más inercia) y la "
              "misma fuerza restauradora del resorte, la aceleración es menor y cada oscilación tarda más. Según "
              "el marco teórico, T = 2π·√(m/k)."),
        ("p", "El periodo al cuadrado dependerá **linealmente** de la masa: T^{2} = (4π^{2}/k)·m. El gráfico de T^{2} "
              "versus m será una recta de pendiente 4π^{2}/k, que permite obtener la constante elástica por un "
              "segundo método."),
    ])
    t2 = tabla(54)
    inner2 = t2.cell(0, 0).tables[0]
    for i, (m, t, T, T2) in enumerate(zip(p2["masas_kg"], p2["t10_s"], p2["T_s"], p2["T2_s2"])):
        fila = inner2.rows[i + 1]
        for j, txt in enumerate([coma(m, 3), coma(t, 2), coma(T, 3), coma(T2, 4)]):
            celda_simple(fila.cells[j + 1], txt, tam=9)
    agregar_a_celda(t2.cell(0, 0), [("p", "t: tiempo de 10 oscilaciones; T = t/10.")])
    agregar_a_celda(t2.cell(0, 1), [("img", g2), ("eq", f"T^{{2}} = {coma(a2, 4)}·m{sg(b2, 4)}   (T^{{2}} en s^{{2}}, m en kg)")],
                    ancho_img_cm=8)
    agregar_a_celda(tabla(58).cell(0, 0), [
        ("p", "De T^{2} = (4π^{2}/k)·m, la pendiente del gráfico 2 es 4π^{2}/k. Por lo tanto:"),
        ("eq", f"k = 4π^{{2}}/pendiente = 4π^{{2}}/({coma(a2, 4)} s^{{2}}/kg) = **{coma(k2, 1)} N/m**"),
        ("p", "(Unidades: kg/s^{2} = N/m.)"),
    ])

    # Complete la frase
    respuestas = ["disminuye.", "aumenta.", "aumenta.", "disminuye."]
    k_resp = 0
    for el in cuerpo:
        if el.tag != qn("w:p") or k_resp >= 4:
            continue
        ts = list(el.iter(qn("w:t")))
        if any("___" in (t.text or "") for t in ts):
            for t in ts:
                if "___" in (t.text or ""):
                    t.text = re.sub(r"_+", respuestas[k_resp], t.text, count=1)
                    break
            k_resp += 1
    assert k_resp == 4, "no se encontraron las 4 líneas para completar"

    agregar_a_celda(tabla(75).cell(0, 0), [
        ("p", f"Los datos del gráfico 2 son acordes a lo esperado: T^{{2}} aumenta linealmente con la masa "
              f"(R^{{2}} = {coma(r22, 4)}), como predice T^{{2}} = (4π^{{2}}/k)·m, lo que confirma la hipótesis de la parte 2. "
              f"La relación algebraica obtenida es T^{{2}} = {coma(a2, 4)}·m{sg(b2, 4)}."),
        ("p", f"La constante elástica obtenida por el método estático (parte 1) es k_{{1}} = {coma(k1, 1)} N/m y por el "
              f"método dinámico (parte 2) es k_{{2}} = {coma(k2, 1)} N/m. La diferencia porcentual es "
              f"|k_{{1}} − k_{{2}}| / k_{{prom}} · 100 = {coma(dif, 1)} %, "
              + ("menor que 10 %, por lo que ambos métodos son consistentes y describen el mismo resorte."
                 if dif <= 10 else "mayor que 10 %, una diferencia apreciable entre los métodos.")),
        ("p", f"El intercepto positivo del gráfico 2 ({coma(b2, 4)} s^{{2}}) se explica porque el resorte también "
              "oscila y aporta parte de su masa (aproximadamente un tercio de ella) al sistema: T^{2} = (4π^{2}/k)·(m + m_{ef}). "
              f"De los datos, m_{{ef}} = intercepto/pendiente ≈ {coma(mef * 1000, 0)} g. Por eso el ajuste lineal con "
              "intercepto da un k más correcto que calcular k punto a punto. Las diferencias restantes se deben al "
              "tiempo de reacción al usar el cronómetro (por eso se midieron 10 oscilaciones) y a la lectura de la "
              "regla."),
    ])

    d.save(salida)
    print("Generado:", salida)
    print("Generado:", g1)
    print("Generado:", g2)
    vm.imprimir(r)


if __name__ == "__main__":
    main()
