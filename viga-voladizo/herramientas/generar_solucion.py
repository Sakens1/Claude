#!/usr/bin/env python3
"""Genera la solución desarrollada del Taller 3 (viga en voladizo) sobre la plantilla original,
con los datos de cualquier laboratorio o grupo.

Cada laboratorio toma sus propios datos, así que el script recibe b, h, L, masas y flechas
y redacta las respuestas (cálculos, comparación, conclusión) según los resultados que obtiene.
Sin argumentos usa un conjunto de datos de ejemplo (regla de acero).

Ejemplos:
    # datos de ejemplo
    python3 generar_solucion.py

    # datos de un laboratorio (dimensiones en mm, masas en g, flechas en mm)
    python3 generar_solucion.py --b 25.1 --h 0.79 --L 300 --dimensiones-en-mm \
        --masas 20 40 60 80 100 --masas-en-g --flechas 6 12.5 18 24.5 30 --flechas-en-mm \
        --etiqueta "Lab martes 10:00, grupo 3" --salida ../solucion/lab_martes_g3.docx

Crea el .docx indicado en --salida y el gráfico .png con el mismo nombre.
Requiere: python-docx, matplotlib.
"""
import argparse
import os
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from docx.table import Table

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import verificar_calculos as vc  # noqa: E402

RAIZ = os.path.dirname(AQUI)
ENUNCIADO = os.path.join(RAIZ, "fuentes", "Taller3_Viga_Voladiza_IN1090C_enunciado.docx")
SALIDA_POR_DEFECTO = os.path.join(RAIZ, "solucion", "Taller3_Viga_Voladiza_RESUELTO.docx")

# ---------------------------------------------------------------- datos de ejemplo
EJEMPLO = dict(b=25.0, h=0.80, L=250.0, dimensiones_en_mm=True,
               masas=[10, 20, 30, 40, 50, 60, 70, 80, 90, 100], masas_en_g=True,
               flechas=[2.5, 5.0, 7.5, 10.0, 12.0, 15.0, 17.5, 20.0, 22.5, 24.5], flechas_en_mm=True)
FILAS_TABLA2 = 10  # filas disponibles en la Tabla 2 de la plantilla

NOMBRE_MATERIAL = {"acero": "acero", "acero_inoxidable": "acero inoxidable", "aluminio": "aluminio",
                   "laton": "latón", "cobre": "cobre"}


def coma(x, dec):
    return f"{x:.{dec}f}".replace(".", ",")


def sig(x, n=3):
    """Cifras significativas manteniendo ceros finales: 0.0008 -> '0,000800'."""
    return format(x, f"#.{n}g").rstrip(".").replace(".", ",")


def cient(x, dec=2):
    """1.0667e-12 -> '1,07×10^{-12}' (con marcado para superíndice)."""
    m, e = f"{x:.{dec}e}".split("e")
    return f"{m.replace('.', ',')}×10^{{{int(e)}}}"


def calcular(a):
    return vc.analizar(a)


# ---------------------------------------------------------------- gráfico
def graficar(r, ruta_png, material):
    F = r["entradas_SI"]["fuerzas_N"]
    y = r["entradas_SI"]["flechas_m"]
    L = r["entradas_SI"]["L_m"]
    m = r["ajuste"]["pendiente_m_por_N"]
    n = r["ajuste"]["intercepto_m"]
    r2 = r["ajuste"]["R2"]
    fig, ax = plt.subplots(figsize=(7.0, 4.4), dpi=200)
    ax.scatter(F, y, color="#1f4e79", zorder=3, label="Datos experimentales (Tabla 2)")
    xs = [0, max(F) * 1.05]
    ax.plot(xs, [m * v + n for v in xs], color="#c00000", lw=1.4,
            label="Línea de tendencia (mínimos cuadrados)")
    ax.set_xlabel("Fuerza aplicada  $F_i$  [N]")
    ax.set_ylabel("Flecha de flexión  $y_{iF}$  [m]")
    ax.set_title(f"Flecha de flexión vs fuerza aplicada (regla de {material}, L = {sig(L)} m)")
    fmt = FuncFormatter(lambda v, _: f"{v:.4f}".rstrip("0").rstrip(".").replace(".", ",") if v else "0")
    ax.xaxis.set_major_formatter(fmt)
    ax.yaxis.set_major_formatter(fmt)
    ax.set_xlim(0, max(F) * 1.05)
    ax.set_ylim(min(0, min(y)), max(y) * 1.12)
    ax.grid(True, ls=":", alpha=0.6)

    def mt(x, dec):  # coma decimal sin espacio extra en mathtext
        return coma(x, dec).replace(",", "{,}")
    signo = "+" if n >= 0 else "-"
    eq = f"$y_F = {mt(m, 5)}\\;F {signo} {mt(abs(n), 5)}$\n$R^2 = {mt(r2, 4)}$"
    ax.text(0.04, 0.95, eq, transform=ax.transAxes, va="top", fontsize=11,
            bbox=dict(boxstyle="round", fc="white", ec="#888888"))
    ax.legend(loc="lower right", fontsize=9)
    fig.tight_layout()
    fig.savefig(ruta_png)
    plt.close(fig)


# ---------------------------------------------------------------- utilidades docx
TOKEN = re.compile(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})")


def agregar_texto(p, texto, tam=10, negrita=False, color=None):
    """Agrega runs a p. Marcado: **negrita**, _{subíndice}, ^{superíndice}."""
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        run_neg, sub, sup = negrita, False, False
        if parte.startswith("**"):
            agregar_texto(p, parte[2:-2], tam=tam, negrita=True, color=color)
            continue
        if parte.startswith("_{"):
            parte, sub = parte[2:-1], True
        elif parte.startswith("^{"):
            parte, sup = parte[2:-1], True
        run = p.add_run(parte)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        run.font.size = Pt(tam)
        run.bold = run_neg
        run.font.subscript = sub
        run.font.superscript = sup
        if color:
            run.font.color.rgb = docx.shared.RGBColor.from_string(color)


def vaciar_celda(celda):
    for p in list(celda.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = celda.paragraphs[0]
    for child in list(p._element):
        if child.tag != qn("w:pPr"):
            p._element.remove(child)
    return p


def llenar_cuadro(celda, bloques):
    """bloques: lista de (tipo, contenido). tipo in {'p','b','eq','img','lista'}."""
    p = vaciar_celda(celda)
    primero = True
    for tipo, cont in bloques:
        if not primero:
            p = celda.add_paragraph()
        primero = False
        p.paragraph_format.space_after = Pt(4)
        if tipo == "p":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            agregar_texto(p, cont)
        elif tipo == "b":
            agregar_texto(p, cont, negrita=True)
        elif tipo == "eq":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            agregar_texto(p, cont, tam=11)
        elif tipo == "lista":
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            agregar_texto(p, "• " + cont)
        elif tipo == "img":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(cont, width=Cm(15))


def celda_simple(celda, texto, centrado=True):
    p = vaciar_celda(celda)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    agregar_texto(p, texto)


# ---------------------------------------------------------------- contenido
def respuestas(r, a, ruta_png, material):
    e = r["entradas_SI"]
    b, h, L = e["b_m"], e["h_m"], e["L_m"]
    I = r["I_m4"]
    m = r["ajuste"]["pendiente_m_por_N"]
    n = r["ajuste"]["intercepto_m"]
    r2 = r["ajuste"]["R2"]
    Y = r["Y_experimental_Pa"]
    Yt = r["Y_teorico_Pa"]
    rango = r["rango_tabla_Pa"]
    err = r["error_porcentual"]
    dY = r["incertidumbre_relativa_Y"]
    contrib = r["contribuciones_incertidumbre"]
    den = 3 * m * I
    razon_I = (b / h) ** 2
    ymax = max(e["flechas_m"])
    pct_flecha = 100 * ymax / L
    npts = len(e["flechas_m"])

    # juicios que dependen de los datos
    lineal = r2 >= 0.98
    intercepto_chico = abs(n) <= max(a.res_y, 0.05 * ymax)
    orden_ok = 1e10 <= Y <= 5e11
    en_rango = bool(rango) and rango[0] <= Y <= rango[1]
    dentro_inc = abs(Y - Yt) <= Y * dY
    subestima = Y < Yt
    Ygpa, Ytgpa = Y / 1e9, Yt / 1e9
    txt_teo = (f"{cient(Yt, 1)} Pa = {coma(Ytgpa, 0)} GPa"
               + (f"; rango típico {coma(rango[0]/1e9, 0)} a {coma(rango[1]/1e9, 0)} GPa" if rango else ""))
    dh = 100 * contrib["3·Δh/h"]
    dL = 100 * contrib["3·ΔL/L"]

    R = {}
    R["1.1"] = [
        ("b", "Hipótesis"),
        ("p", "Si se aplica una fuerza vertical concentrada F en el extremo libre de la regla metálica "
              "empotrada, y se trabaja en el régimen elástico con pequeñas deformaciones (pendientes pequeñas), "
              "entonces la flecha de flexión y_{F} será directamente proporcional a la fuerza aplicada:"),
        ("eq", "y_{F} = c_{f}·F,   con   c_{f} = L^{3} / (3·Y·I)"),
        ("p", "Por lo tanto, el gráfico y_{F} vs F será una recta que pasa aproximadamente por el origen, "
              "cuya pendiente es la constante de flexibilidad c_{f}. Conocidos L e I = b·h^{3}/12, a partir de "
              "esa pendiente se podrá determinar el módulo de Young de la regla, que debería ser cercano al valor "
              f"tabulado para el {material} (Y ≈ {txt_teo.split(';')[0]})."),
    ]
    R["1.4"] = [
        ("p", f"Datos de la Tabla 1: ancho b = {sig(b)} m, espesor h = {sig(h)} m y longitud libre "
              f"L = {sig(L)} m (medida desde el borde del empotramiento hasta el punto donde cuelga la masa)."),
        ("p", "La carga actúa en la dirección del espesor h, por lo que la flexión ocurre en torno al eje neutro "
              "paralelo al ancho b. El momento de inercia de la sección rectangular respecto a ese eje es:"),
        ("eq", "I = (1/12)·b·h^{3}"),
        ("eq", f"I = (1/12)·({sig(b)} m)·({sig(h)} m)^{{3}} = (1/12)·({sig(b)})·({cient(h**3, 3)}) m^{{4}}"),
        ("eq", f"**I = {cient(I, 2)} m^{{4}}**"),
    ]
    signo = "+" if n >= 0 else "−"
    R["1.5"] = [
        ("img", ruta_png),
        ("p", "Ecuación de la línea de tendencia (ajuste lineal por mínimos cuadrados, F en N e y_{F} en m):"),
        ("eq", f"**y_{{F}} = {coma(m, 5)}·F {signo} {coma(abs(n), 5)}      R^{{2}} = {coma(r2, 4)}**"),
    ]
    if intercepto_chico:
        txt_n = (f"Su valor ({coma(n*1000, 2)} mm) es pequeño frente a las flechas medidas y del orden de la "
                 "resolución de la lectura, por lo que se atribuye a errores de lectura y no a un efecto físico.")
    else:
        txt_n = (f"Su valor ({coma(n*1000, 2)} mm) no es despreciable: indica un error de cero (referencia "
                 "sin carga mal tomada, holgura en la mordaza) o que algunos puntos se apartan de la recta.")
    if r2 >= 0.99:
        txt_r2 = "muy cercano a 1, confirma la relación lineal (comportamiento elástico, ley de Hooke) entre la flecha y la fuerza."
    elif lineal:
        txt_r2 = "cercano a 1, indica una relación lineal entre la flecha y la fuerza, con algo de dispersión experimental."
    else:
        txt_r2 = ("alejado de 1: los datos tienen mucha dispersión o no son lineales (posibles grandes "
                  "deflexiones o errores de lectura), por lo que la pendiente es menos confiable.")
    R["1.5.1"] = [
        ("p", "La línea de tendencia tiene la forma y_{F} = m·F + n, que se compara con el modelo teórico "
              "y_{F} = c_{f}·F:"),
        ("lista", f"**Pendiente m = {coma(m, 5)} m/N**: corresponde a la constante de flexibilidad "
                  "c_{f} = L^{3}/(3·Y·I). Indica cuántos metros desciende el extremo libre por cada newton aplicado. "
                  "Se asocia a la longitud L (geometría), al momento de inercia I (forma y tamaño de la sección: b y h) "
                  "y al módulo de Young Y (material). Su inverso, k = 3YI/L^{3} ≈ "
                  f"{coma(1/m, 1)} N/m, es la rigidez de la viga."),
        ("lista", f"**Intercepto n = {coma(n, 5)} m**: es la flecha cuando F = 0. Teóricamente es cero, "
                  "porque la referencia se tomó en el extremo libre sin carga. " + txt_n),
        ("lista", f"**R^{{2}} = {coma(r2, 4)}**: " + txt_r2),
    ]
    R["1.5.2"] = [
        ("p", "Igualando la pendiente experimental con la constante de flexibilidad y despejando Y:"),
        ("eq", "m = c_{f} = L^{3}/(3·Y·I)   ⟹   Y = L^{3}/(3·m·I)"),
        ("eq", f"Y = ({sig(L)} m)^{{3}} / (3 · {coma(m, 5)} m/N · {cient(I, 3)} m^{{4}})"),
        ("eq", f"Y = {cient(L**3, 4)} m^{{3}} / {cient(den, 3)} m^{{5}}/N"),
        ("eq", f"**Y_{{exp}} = {cient(Y, 2)} Pa ≈ {coma(Ygpa, 0)} GPa**"),
        ("p", f"Estimación de incertidumbre (propagación): ΔY/Y = 3ΔL/L + 3Δh/h + Δb/b + Δm/m ≈ "
              f"{coma(100*dY, 1)} %, es decir Y_{{exp}} = ({coma(Ygpa, 0)} ± {coma(Ygpa*dY, 0)}) GPa. "
              + ("El término dominante es el del espesor " if dh >= dL else "El término dominante es el de la longitud ")
              + f"(3Δh/h ≈ {coma(dh, 1)} %, 3ΔL/L ≈ {coma(dL, 1)} %), porque h y L aparecen al cubo."),
    ]
    if not orden_ok:
        juicio = ("El valor experimental **no tiene el orden de magnitud esperado** para un metal (10^{10} a "
                  "10^{11} Pa). Esto indica un error en los datos (unidades, espesor h mal medido o b y h "
                  "intercambiados) más que una propiedad del material; deben revisarse las mediciones.")
    else:
        juicio = "El valor experimental tiene el orden de magnitud correcto (10^{11} Pa)"
        if rango:
            juicio += (f", queda dentro del rango tabulado del {material}" if en_rango
                       else f", pero queda fuera del rango tabulado del {material}")
        juicio += (" y el valor teórico cae dentro del intervalo de incertidumbre experimental, por lo que la "
                   "diferencia es compatible con los errores de medición." if dentro_inc else
                   ". El valor teórico queda fuera del intervalo de incertidumbre estimado, lo que sugiere un "
                   "error sistemático además de los errores aleatorios.")
    if subestima:
        sesgo = ("Empotramiento no perfectamente rígido: la mordaza permite un pequeño giro que aumenta la flecha "
                 "y hace subestimar Y, coherente con obtener un valor menor que el teórico.")
    else:
        sesgo = ("El valor obtenido es mayor que el teórico: Y se sobreestima si L se midió más larga que la "
                 "longitud efectiva (punto de carga) o si h se midió menor que el espesor real.")
    R["1.6"] = [
        ("p", f"Valor teórico: módulo de Young del {material} Y_{{teo}} = {txt_teo} (tablas de propiedades de "
              "materiales, p. ej. Serway & Jewett, Física para ciencias e ingeniería)."),
        ("eq", "E% = |Y_{exp} − Y_{teo}| / Y_{teo} × 100"),
        ("eq", f"E% = |{coma(Ygpa, 1)} − {coma(Ytgpa, 0)}| / {coma(Ytgpa, 0)} × 100 = **{coma(err, 1)} %**"),
        ("p", juicio),
        ("p", "Fuentes de error que explican la diferencia:"),
        ("lista", f"Medición del espesor h: su incertidumbre aporta ≈ {coma(dh, 1)} % de error en Y "
                  "(dependencia h^{3})."),
        ("lista", "Definición de L: debe medirse desde el borde de la mordaza hasta el punto de aplicación de la "
                  f"carga; su incertidumbre aporta ≈ {coma(dL, 1)} % en Y (dependencia L^{{3}})."),
        ("lista", sesgo),
        ("lista", "Lectura de la flecha con error de paralaje y resolución limitada; oscilación de la masa."),
        ("lista", "Masa del portapesas y la aproximación de pendientes pequeñas "
                  f"(flecha máxima ≈ {coma(pct_flecha, 0)} % de L)."),
    ]
    R["1.7.1"] = [
        ("p", "Las vigas pueden sufrir deformaciones por:"),
        ("lista", "**Flexión**: curvatura del eje de la viga por cargas perpendiculares a él (momento flector)."),
        ("lista", "**Corte o cizalle**: deslizamiento relativo entre secciones por el esfuerzo cortante."),
        ("lista", "**Axial (tracción o compresión)**: cuando las cargas tienen componente paralela al eje; en "
                  "compresión puede producirse pandeo."),
        ("lista", "**Torsión**: giro de las secciones cuando la carga no pasa por el eje de la viga."),
        ("p", "La más común es la **flexión** (acompañada de esfuerzo cortante), porque en la mayoría de los casos "
              "las cargas, como el peso, actúan perpendiculares al eje de la viga. Es la deformación estudiada en "
              "este taller."),
    ]
    R["1.7.2"] = [
        ("lista", "En posición horizontal las cargas gravitacionales actúan perpendiculares al eje, de modo que la "
                  "viga trabaja a flexión y transmite esas cargas a los apoyos o columnas."),
        ("lista", "Permiten salvar luces (distancia entre apoyos) y generar espacios libres: losas, techos, puentes, "
                  "y en voladizo, balcones, aleros, graderías, estanterías cantilever y puentes construidos por avance "
                  "en voladizo, sin apoyos intermedios."),
        ("lista", "Orientando la sección con su mayor dimensión en dirección vertical (en la dirección de la carga) "
                  "se maximiza I ∝ h^{3} y la rigidez sin agregar material. En la regla del taller, si se cargara "
                  f"de canto (b en dirección de la carga) I aumentaría en un factor (b/h)^{{2}} ≈ {coma(razon_I, 0)}: "
                  "por eso las vigas reales son altas y angostas (o con perfil I)."),
    ]

    if orden_ok and err <= 10:
        c_ra = (f"Esto permite identificar el material de la regla como {material}." if en_rango or not rango
                else f"El resultado es cercano al del {material}, aunque fuera de su rango tabulado.")
    else:
        c_ra = (f"La diferencia es grande, por lo que con estos datos no se puede confirmar que la regla sea de "
                f"{material}; es necesario revisar las mediciones (especialmente h y L) y las unidades.")
    if lineal and intercepto_chico:
        c_hip = (f"La hipótesis se confirma: la flecha de flexión aumenta linealmente con la fuerza aplicada "
                 f"(R^{{2}} = {coma(r2, 4)}) y el intercepto es prácticamente nulo ({coma(n*1000, 2)} mm), tal "
                 "como predice y_{F} = c_{f}·F.")
    elif lineal:
        c_hip = (f"La hipótesis se confirma en cuanto a la linealidad (R^{{2}} = {coma(r2, 4)}), aunque el "
                 f"intercepto ({coma(n*1000, 2)} mm) indica un error de cero en la referencia de las flechas.")
    else:
        c_hip = (f"La hipótesis no se confirma plenamente: el ajuste lineal es pobre (R^{{2}} = {coma(r2, 4)}), "
                 "por dispersión de las lecturas o por salir del régimen de pequeñas deformaciones.")
    c_hip += (" La pendiente de la recta corresponde a la constante de flexibilidad "
              f"c_{{f}} = L^{{3}}/(3YI) = {coma(m, 4)} m/N.")
    c_err = ("Del análisis de datos se concluye que las mediciones más críticas son el espesor h y la longitud L, "
             f"por su dependencia cúbica; la incertidumbre estimada es ≈ {coma(100*dY, 0)} % y el error obtenido "
             f"es {coma(err, 1)} %"
             + (", por lo que la diferencia es atribuible a errores experimentales. " if dentro_inc else
                ", mayor que la incertidumbre, lo que apunta a un error sistemático. ")
             + "Para mejorar el resultado se propone medir h con micrómetro en varios puntos y promediar, asegurar "
               "firmemente el empotramiento y medir la flecha evitando el paralaje.")
    c_mod = ("Finalmente, el modelo de viga en voladizo es válido dentro de sus supuestos (material homogéneo y "
             "elástico, sección uniforme, peso propio despreciable y pequeñas deformaciones)"
             + (f", los cuales se cumplieron al mantener la flecha máxima en ≈ {coma(pct_flecha, 0)} % de L."
                if pct_flecha <= 10 else
                f"; la flecha máxima llegó a ≈ {coma(pct_flecha, 0)} % de L, en el límite de la aproximación de "
                "pendientes pequeñas, lo que puede afectar los últimos puntos."))
    if npts < 5:
        c_mod += f" Además, con solo {npts} mediciones el ajuste es poco robusto; se recomiendan al menos 8."
    R["1.8"] = [
        ("p", "Respecto al resultado de aprendizaje, se determinó experimentalmente el módulo de Young de una "
              f"regla metálica, obteniéndose Y_{{exp}} = {cient(Y, 2)} Pa ≈ {coma(Ygpa, 0)} GPa, con un error "
              f"porcentual de {coma(err, 1)} % respecto al valor tabulado del {material} ({coma(Ytgpa, 0)} GPa). "
              + c_ra),
        ("p", c_hip),
        ("p", c_err),
        ("p", c_mod),
    ]
    return R


def decimales(valores, minimo, maximo):
    """Menor cantidad de decimales (entre minimo y maximo) que representa todos los valores."""
    for d in range(minimo, maximo + 1):
        if all(abs(round(v, d) - v) < 1e-12 for v in valores):
            return d
    return maximo


# ---------------------------------------------------------------- armado del documento
def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--b", type=float, help="ancho de la regla")
    p.add_argument("--h", type=float, help="espesor de la regla")
    p.add_argument("--L", type=float, help="longitud libre")
    p.add_argument("--masas", type=float, nargs="+")
    p.add_argument("--flechas", type=float, nargs="+")
    p.add_argument("--masas-en-g", action="store_true")
    p.add_argument("--flechas-en-mm", action="store_true")
    p.add_argument("--dimensiones-en-mm", action="store_true")
    p.add_argument("--g", type=float, default=9.80)
    p.add_argument("--material", default="acero", choices=sorted(vc.TABLA_Y))
    p.add_argument("--Y-teorico", dest="Y_teorico", type=float, help="valor teórico en Pa (anula --material)")
    p.add_argument("--res-L", type=float, default=1e-3)
    p.add_argument("--res-h", type=float, default=1e-5)
    p.add_argument("--res-b", type=float, default=2e-5)
    p.add_argument("--res-y", type=float, default=5e-4, help="resolución de lectura de la flecha en m")
    p.add_argument("--etiqueta", help="texto para el cuadro de integrantes (laboratorio, grupo, fecha)")
    p.add_argument("--salida", default=SALIDA_POR_DEFECTO, help="ruta del .docx a generar")
    a = p.parse_args(argv)

    datos = [a.b, a.h, a.L, a.masas, a.flechas]
    if all(v is None for v in datos):
        for k, v in EJEMPLO.items():
            setattr(a, k, v)
        if not a.etiqueta:
            a.etiqueta = "Datos de ejemplo (regla de acero)"
    elif any(v is None for v in datos):
        p.error("indique --b, --h, --L, --masas y --flechas (o ninguno para usar los datos de ejemplo)")
    if len(a.masas) > FILAS_TABLA2:
        p.error(f"la Tabla 2 de la plantilla tiene {FILAS_TABLA2} filas; se recibieron {len(a.masas)} mediciones")
    for k in ("informado_I", "informado_pendiente", "informado_Y", "informado_error"):
        setattr(a, k, None)

    r = calcular(a)
    material = "material de referencia" if a.Y_teorico else NOMBRE_MATERIAL[a.material]
    salida = os.path.abspath(a.salida)
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    ruta_png = os.path.splitext(salida)[0] + "_grafico.png"
    if salida == os.path.abspath(SALIDA_POR_DEFECTO):
        ruta_png = os.path.join(os.path.dirname(salida), "grafico_yF_vs_F.png")
    graficar(r, ruta_png, material)

    d = docx.Document(ENUNCIADO)
    cuerpo = list(d.element.body.iterchildren())

    def tabla(i):
        el = cuerpo[i]
        assert el.tag == qn("w:tbl"), f"elemento {i} no es tabla"
        return Table(el, d)

    # Integrantes
    llenar_cuadro(tabla(3).cell(0, 0), [("b", "SOLUCIÓN DESARROLLADA / REFERENCIA DOCENTE")])
    llenar_cuadro(tabla(3).cell(0, 1), [("p", a.etiqueta or "")])

    # Tabla 1: dimensiones
    e = r["entradas_SI"]
    t1 = tabla(19)
    celda_simple(t1.cell(0, 1), f"{sig(e['b_m'])} m")
    celda_simple(t1.cell(1, 1), f"{sig(e['h_m'])} m")
    celda_simple(t1.cell(2, 1), f"{sig(e['L_m'])} m")

    # Tabla 2 (dentro de un cuadro de texto: hay 2 copias, AlternateContent Choice y Fallback)
    M, F, Yf = e["masas_kg"], e["fuerzas_N"], e["flechas_m"]
    dm = decimales(M, 3, 5)
    dy = decimales(Yf, 4, 6)
    for tel in cuerpo[32].iter(qn("w:tbl")):
        t2 = Table(tel, d)
        assert len(t2.rows) == FILAS_TABLA2 + 1
        for i in range(len(M)):
            fila = t2.rows[i + 1]
            celda_simple(fila.cells[1], coma(M[i], dm))
            celda_simple(fila.cells[2], coma(F[i], 3 if dm == 3 else 4))
            celda_simple(fila.cells[3], coma(Yf[i], dy))

    # Cuadros de respuesta (tablas 1x1 bajo cada pregunta)
    R = respuestas(r, a, ruta_png, material)
    mapa = {15: ("1.1", "1.1"), 51: ("1.4", "1.4"), 54: ("1.5", "1.5 Construya"),
            57: ("1.5.1", "1.5.1"), 61: ("1.5.2", "1.5.2"), 64: ("1.6", "1.6"),
            68: ("1.7.1", "1.7.1"), 72: ("1.7.2", "1.7.2"), 76: ("1.8", "1.8")}
    for idx, (clave, texto) in mapa.items():
        cerca = "".join(t.text or "" for el in cuerpo[max(0, idx - 4):idx] for t in el.iter(qn("w:t")))
        assert texto in cerca, f"La tabla {idx} no sigue a la pregunta {clave}"
        llenar_cuadro(tabla(idx).cell(0, 0), R[clave])

    d.save(salida)
    print("Generado:", salida)
    print("Generado:", ruta_png)
    vc.imprimir(r)


if __name__ == "__main__":
    main()
