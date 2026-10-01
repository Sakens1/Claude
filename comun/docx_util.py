"""Utilidades para escribir respuestas dentro de las plantillas .docx de los talleres (python-docx).

Marcado admitido en los textos: **negrita**, _{subíndice}, ^{superíndice}.
Bloques: lista de (tipo, contenido) con tipo en {"p", "b", "eq", "lista", "img"}.
"""
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

TOKEN = re.compile(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})")


def coma(x, dec):
    return f"{x:.{dec}f}".replace(".", ",")


def sig(x, n=3):
    """Cifras significativas manteniendo ceros finales: 0.0008 -> '0,000800'."""
    return format(x, f"#.{n}g").rstrip(".").replace(".", ",")


def cient(x, dec=2):
    """1.0667e-12 -> '1,07×10^{-12}' (con marcado para superíndice)."""
    m, e = f"{x:.{dec}e}".split("e")
    return f"{m.replace('.', ',')}×10^{{{int(e)}}}"


def agregar_texto(p, texto, tam=10, negrita=False, color=None):
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        sub = sup = False
        if parte.startswith("**"):
            agregar_texto(p, parte[2:-2], tam=tam, negrita=True, color=color)
            continue
        if parte.startswith("_{"):
            parte, sub = parte[2:-1], True
        elif parte.startswith("^{"):
            parte, sup = parte[2:-1], True
        run = p.add_run(parte)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        run.font.size = Pt(tam)
        run.bold = negrita
        run.font.subscript = sub
        run.font.superscript = sup
        if color:
            run.font.color.rgb = docx.shared.RGBColor.from_string(color)


def vaciar_celda(celda):
    for p in list(celda.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = celda.paragraphs[0]
    for child in list(p._element):
        if child.tag != qn("w:pPr"):
            p._element.remove(child)
    return p


def _escribir_bloque(p, tipo, cont, ancho_img_cm):
    p.paragraph_format.space_after = Pt(4)
    if tipo == "p":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        agregar_texto(p, cont)
    elif tipo == "b":
        agregar_texto(p, cont, negrita=True)
    elif tipo == "eq":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        agregar_texto(p, cont, tam=11)
    elif tipo == "lista":
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        agregar_texto(p, "• " + cont)
    elif tipo == "img":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(cont, width=Cm(ancho_img_cm))


def llenar_cuadro(celda, bloques, ancho_img_cm=15):
    """Reemplaza el contenido de la celda por los bloques."""
    p = vaciar_celda(celda)
    for i, (tipo, cont) in enumerate(bloques):
        if i:
            p = celda.add_paragraph()
        _escribir_bloque(p, tipo, cont, ancho_img_cm)


def agregar_a_celda(celda, bloques, ancho_img_cm=15):
    """Agrega los bloques al final de la celda, conservando lo que ya tiene (p. ej. el enunciado).
    Quita los párrafos vacíos del final antes de agregar."""
    while len(celda.paragraphs) > 1 and not celda.paragraphs[-1].text.strip() \
            and not celda.paragraphs[-1]._element.findall(".//" + qn("w:drawing")):
        el = celda.paragraphs[-1]._element
        el.getparent().remove(el)
    for tipo, cont in bloques:
        _escribir_bloque(celda.add_paragraph(), tipo, cont, ancho_img_cm)


def celda_simple(celda, texto, centrado=True, tam=10):
    p = vaciar_celda(celda)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    agregar_texto(p, texto, tam=tam)
